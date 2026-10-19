from pathlib import Path
import json, hashlib
import openpyxl

root = Path(__file__).resolve().parents[1]
source = root.parent / 'Moduli 215797176 08_10_2026 10.51.Xlsx'
book = openpyxl.load_workbook(source, data_only=False)
cached = openpyxl.load_workbook(source, data_only=True)
result = {'source': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'sheets': []}
for sheet in book:
    cells = []
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                cells.append({'cell': cell.coordinate, 'value': str(cell.value), 'cached': str(cached[sheet.title][cell.coordinate].value or ''), 'bold': cell.font.bold, 'size': cell.font.sz})
    result['sheets'].append({'name': sheet.title, 'print_area': str(sheet.print_area), 'merged': [str(r) for r in sheet.merged_cells.ranges], 'rows': sheet.max_row, 'columns': sheet.max_column, 'cells': cells})
for sheet in result['sheets']:
    print(json.dumps({'sheet':sheet['name'],'cells':len(sheet['cells']),'linked_fields':sum(c['value'].startswith('=') for c in sheet['cells'])},ensure_ascii=False))

mapping = {'D20':'{{compagnia}}','D21':'{{ramo}}','D22':'{{numero_polizza}}','D23':'{{decorrenza}}','D25':'{{titolo_cliente}}','D26':'{{contraente}}','D27':'{{codice_fiscale_piva}}','D28':'{{indirizzo}}','D29':'{{nascita}}','D30':'{{firmatario}}','D31':'{{qualifica_firmatario}}','D33':'{{luogo}}','D34':'{{data_documento}}'}
categories = {'Allegato 3':'annex_3','Allegato 4 Ter':'annex_4_ter','Allegato 4bis':'annex_4_bis','Allegato 4':'annex_4','Consegna Documenti':'delivery_receipt','Dichiarazione Coerenza':'consistency','Consenso Invio Elettronico':'consent'}
templates = []
for sheet in book:
    if sheet.title not in categories:
        continue
    lines = []
    for row in sheet:
        parts = []
        for cell in row:
            value = cell.value
            if value is None:
                continue
            if isinstance(value,str) and value.startswith('=Dati!'):
                reference = value.split('!')[1]
                value = mapping.get(reference)
                if value is None:
                    value = cached['Dati'][reference].value
                    if hasattr(value,'strftime'):
                        value = value.strftime('%d/%m/%Y')
            if value is not None:
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                parts.append(str(value))
        if parts:
            lines.append('   '.join(parts))
    body = '\n\n'.join(lines)
    # The source incorrectly repeats the contractor where the signatory must be identified.
    body = body.replace('Il sottoscritto   {{contraente}}', 'Il sottoscritto   {{firmatario}}')
    # The role label must use the role field, even where the original formula repeats the signatory.
    body = body.replace('In qualità di   {{firmatario}}', 'In qualità di   {{qualifica_firmatario}}')
    legacy = sheet.title.startswith('Allegato') or sheet.title == 'Consegna Documenti'
    templates.append({'title':sheet.title,'category':categories[sheet.title],'scope':'client' if sheet.title=='Consenso Invio Elettronico' else 'policy','body':body,'signature_required':not sheet.title.startswith('Allegato'),'required':False,'status':'legacy' if legacy else 'draft','source_sheet':sheet.title,'source_hash':result['sha256'],'source_warning':'Il testo proviene dal precedente gestionale. Verificare testi, dati dell’intermediario e campi prima dell’uso.' + (' Contiene riferimenti ai precedenti allegati IVASS; la ricevuta richiama anche il precedente portale. Occorre una versione aggiornata.' if legacy else '')})
target = root/'resources/insurance/forms-source.json'
target.parent.mkdir(parents=True,exist_ok=True)
target.write_text(json.dumps(templates,ensure_ascii=False,indent=2),encoding='utf-8')
